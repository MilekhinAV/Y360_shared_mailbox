#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Массовое создание общих ящиков Яндекс 360 и назначение прав из XLSX/CSV."""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import time
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


VALID_ROLES = {
    "shared_mailbox_imap_admin",
    "shared_mailbox_sender",
    "shared_mailbox_half_sender",
    "shared_mailbox_owner",
}

ROLE_ALIASES = {
    "imap": "shared_mailbox_imap_admin",
    "imap_admin": "shared_mailbox_imap_admin",
    "sender": "shared_mailbox_sender",
    "send": "shared_mailbox_sender",
    "half_sender": "shared_mailbox_half_sender",
    "half-sender": "shared_mailbox_half_sender",
    "owner": "shared_mailbox_owner",
    "full": "shared_mailbox_owner",
    "all": "shared_mailbox_owner",
}

REQUIRED_HEADERS = {"UID", "email", "name", "description", "roles"}
ROLE_SPLIT_RE = re.compile(r"[,;\s]+")


class ValidationError(Exception):
    pass


@dataclass
class AccessGrant:
    row_number: int
    identifier: str
    roles: list[str]
    actor_id: Optional[str] = None


@dataclass
class MailboxPlan:
    email: str
    name: str
    description: str
    first_row: int
    grants: list[AccessGrant] = field(default_factory=list)
    resource_id: Optional[str] = None


def die(message: str, code: int = 1) -> None:
    print(f"ОШИБКА: {message}", file=sys.stderr)
    raise SystemExit(code)


def load_dotenv_file(path: Path = Path(".env")) -> None:
    """Загружает простой KEY=VALUE без перезаписи переменных окружения."""
    if not path.is_file():
        return
    for line_number, raw in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[7:].lstrip()
        if "=" not in line:
            raise ValidationError(f".env, строка {line_number}: ожидается KEY=VALUE")
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {'"', "'"}:
            value = value[1:-1]
        if key:
            os.environ.setdefault(key, value)


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_header(value: object) -> str:
    header = cell_text(value).lstrip("\ufeff")
    return "UID" if header.casefold() == "uid" else header.casefold()


def parse_roles(raw: str, row_number: int) -> list[str]:
    tokens = [part.casefold() for part in ROLE_SPLIT_RE.split(raw.strip()) if part]
    if not tokens:
        raise ValidationError(f"строка {row_number}: поле roles обязательно")

    roles: list[str] = []
    unknown: list[str] = []
    for token in tokens:
        role = ROLE_ALIASES.get(token, token)
        if role not in VALID_ROLES:
            unknown.append(token)
        elif role not in roles:
            roles.append(role)

    if unknown:
        allowed = ", ".join(sorted(VALID_ROLES | set(ROLE_ALIASES)))
        raise ValidationError(
            f"строка {row_number}: неизвестные роли: {', '.join(unknown)}. "
            f"Допустимые значения: {allowed}"
        )

    if "shared_mailbox_owner" in roles:
        return ["shared_mailbox_owner"]
    return roles


def iter_xlsx_rows(path: Path, sheet_name: str) -> Iterable[tuple[int, dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        raise ValidationError(
            f"в файле нет листа '{sheet_name}'. Доступные листы: {available}"
        )
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration:
        raise ValidationError("таблица пуста")

    headers = [normalize_header(value) for value in header_values]
    validate_headers(headers)
    for row_number, values in enumerate(rows, start=2):
        row = {headers[i]: cell_text(values[i]) if i < len(values) else "" for i in range(len(headers))}
        if any(row.values()):
            yield row_number, row


def detect_csv_dialect(path: Path) -> csv.Dialect:
    sample = path.read_text(encoding="utf-8-sig")[:8192]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        return csv.excel


def iter_csv_rows(path: Path) -> Iterable[tuple[int, dict[str, str]]]:
    dialect = detect_csv_dialect(path)
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.reader(source, dialect)
        try:
            header_values = next(reader)
        except StopIteration:
            raise ValidationError("таблица пуста")
        headers = [normalize_header(value) for value in header_values]
        validate_headers(headers)
        for row_number, values in enumerate(reader, start=2):
            row = {headers[i]: cell_text(values[i]) if i < len(values) else "" for i in range(len(headers))}
            if any(row.values()):
                yield row_number, row


def validate_headers(headers: list[str]) -> None:
    actual = set(headers)
    missing = REQUIRED_HEADERS - actual
    if missing:
        raise ValidationError(
            "отсутствуют обязательные столбцы: " + ", ".join(sorted(missing))
        )
    if len(headers) != len(set(headers)):
        raise ValidationError("в таблице есть повторяющиеся заголовки")


def load_plan(path: Path, sheet_name: str) -> list[MailboxPlan]:
    if not path.is_file():
        raise ValidationError(f"файл не найден: {path}")

    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        rows = iter_xlsx_rows(path, sheet_name)
    elif suffix in {".csv", ".tsv"}:
        rows = iter_csv_rows(path)
    else:
        raise ValidationError("поддерживаются только файлы .xlsx, .csv и .tsv")

    plans: OrderedDict[str, MailboxPlan] = OrderedDict()
    errors: list[str] = []

    for row_number, row in rows:
        mailbox_email = row["email"].strip()
        mailbox_name = row["name"].strip()
        description = row["description"].strip()
        identifier = row["UID"].strip()

        if not mailbox_email:
            errors.append(f"строка {row_number}: поле email обязательно")
        elif any(char.isspace() for char in mailbox_email):
            errors.append(f"строка {row_number}: email содержит пробелы")
        if not mailbox_name:
            errors.append(f"строка {row_number}: поле name обязательно")
        if not identifier:
            errors.append(f"строка {row_number}: поле UID обязательно")

        try:
            roles = parse_roles(row["roles"], row_number)
        except ValidationError as exc:
            errors.append(str(exc))
            roles = []

        if not mailbox_email or not mailbox_name or not identifier or not roles:
            continue

        key = mailbox_email.casefold()
        plan = plans.get(key)
        if plan is None:
            plan = MailboxPlan(mailbox_email, mailbox_name, description, row_number)
            plans[key] = plan
        elif plan.name != mailbox_name or plan.description != description:
            errors.append(
                f"строка {row_number}: для ящика '{mailbox_email}' name/description "
                f"отличаются от строки {plan.first_row}"
            )
            continue

        duplicate = next(
            (grant for grant in plan.grants if grant.identifier.casefold() == identifier.casefold()),
            None,
        )
        if duplicate:
            errors.append(
                f"строка {row_number}: повторное назначение '{identifier}' для "
                f"ящика '{mailbox_email}' (первая строка {duplicate.row_number})"
            )
            continue
        plan.grants.append(AccessGrant(row_number, identifier, roles))

    if errors:
        raise ValidationError("\n".join(errors))
    if not plans:
        raise ValidationError("таблица не содержит строк с данными")
    return list(plans.values())


class Api360:
    def __init__(self, token: str, org_id: str, timeout: float = 30.0):
        self.org_id = org_id
        self.timeout = timeout
        self.base_admin = f"https://api360.yandex.net/admin/v1/org/{org_id}"
        self.base_directory = f"https://api360.yandex.net/directory/v1/org/{org_id}"
        self.headers = {
            "Authorization": f"OAuth {token}",
            "Accept": "application/json",
            "Content-Type": "application/json",
            "User-Agent": "Y360-shared-mailbox-bulk/2.0",
        }

    def request(
        self,
        method: str,
        url: str,
        *,
        params: Optional[dict] = None,
        body: Optional[dict] = None,
    ) -> dict:
        if params:
            url = f"{url}?{urlencode(params)}"
        payload = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")

        for attempt in range(5):
            request = Request(url, data=payload, headers=self.headers, method=method)
            try:
                with urlopen(request, timeout=self.timeout) as response:
                    raw = response.read()
                    return json.loads(raw.decode("utf-8")) if raw else {}
            except HTTPError as exc:
                raw = exc.read().decode("utf-8", errors="replace")
                retryable = exc.code in {429, 500, 502, 503, 504}
                if retryable and attempt < 4:
                    retry_after = exc.headers.get("Retry-After", "")
                    delay = float(retry_after) if retry_after.isdigit() else 2**attempt
                    time.sleep(delay)
                    continue
                try:
                    details = json.loads(raw)
                except json.JSONDecodeError:
                    details = raw
                raise RuntimeError(f"HTTP {exc.code}: {details}") from exc
            except URLError as exc:
                if attempt < 4:
                    time.sleep(2**attempt)
                    continue
                raise RuntimeError(f"ошибка соединения: {exc.reason}") from exc
        raise RuntimeError("не удалось выполнить запрос после повторных попыток")

    def list_users(self) -> list[dict]:
        users: list[dict] = []
        page = 1
        while True:
            data = self.request(
                "GET",
                f"{self.base_directory}/users",
                params={"page": page, "perPage": 1000},
            )
            users.extend(data.get("users", []))
            pages = int(data.get("pages") or 1)
            if page >= pages:
                return users
            page += 1

    def create_shared_mailbox(self, plan: MailboxPlan) -> str:
        body = {"email": plan.email, "name": plan.name}
        if plan.description:
            body["description"] = plan.description
        data = self.request("PUT", f"{self.base_admin}/mailboxes/shared", body=body)
        resource_id = data.get("resourceId")
        if not resource_id:
            raise RuntimeError(f"API не вернул resourceId: {data}")
        return str(resource_id)

    def set_access(self, resource_id: str, grant: AccessGrant, notify: str) -> str:
        data = self.request(
            "POST",
            f"{self.base_admin}/mailboxes/set/{resource_id}",
            params={"actorId": grant.actor_id, "notify": notify},
            body={"roles": grant.roles},
        )
        task_id = data.get("taskId")
        if not task_id:
            raise RuntimeError(f"API не вернул taskId: {data}")
        return str(task_id)

    def wait_task(self, task_id: str, wait_seconds: float = 60.0) -> str:
        deadline = time.monotonic() + wait_seconds
        while True:
            data = self.request(
                "GET", f"{self.base_admin}/mailboxes/tasks/{task_id}"
            )
            status = str(data.get("status", "")).casefold()
            if status in {"complete", "error"}:
                return status
            if time.monotonic() >= deadline:
                return "timeout"
            time.sleep(1)


def build_user_index(users: list[dict]) -> tuple[dict[str, str], list[str]]:
    candidates: dict[str, set[str]] = {}
    warnings: list[str] = []
    for user in users:
        user_id = cell_text(user.get("id"))
        if not user_id:
            continue
        keys = [cell_text(user.get("email")), cell_text(user.get("nickname"))]
        for alias in user.get("aliases") or []:
            keys.append(cell_text(alias))
        for key in keys:
            if key:
                candidates.setdefault(key.casefold(), set()).add(user_id)

    index: dict[str, str] = {}
    for key, ids in candidates.items():
        if len(ids) == 1:
            index[key] = next(iter(ids))
        else:
            warnings.append(f"идентификатор '{key}' соответствует нескольким UID")
    return index, warnings


def resolve_actor_ids(plans: list[MailboxPlan], api: Api360) -> list[str]:
    unresolved = [
        grant
        for plan in plans
        for grant in plan.grants
        if not grant.identifier.isdigit()
    ]
    if not unresolved:
        for plan in plans:
            for grant in plan.grants:
                grant.actor_id = grant.identifier
        return []

    users = api.list_users()
    index, warnings = build_user_index(users)
    errors: list[str] = []
    for plan in plans:
        for grant in plan.grants:
            if grant.identifier.isdigit():
                grant.actor_id = grant.identifier
                continue
            grant.actor_id = index.get(grant.identifier.casefold())
            if not grant.actor_id:
                errors.append(
                    f"строка {grant.row_number}: сотрудник '{grant.identifier}' не найден"
                )
    if errors:
        raise ValidationError("\n".join(errors))
    return warnings


def write_report(path: Path, rows: list[dict[str, str]]) -> None:
    fieldnames = [
        "mailbox_email",
        "resource_id",
        "delegate",
        "actor_id",
        "roles",
        "status",
        "details",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Массовое создание общих ящиков Яндекс 360 из XLSX/CSV"
    )
    parser.add_argument("file", help="Путь к .xlsx, .csv или .tsv")
    parser.add_argument(
        "--sheet", default="main_list", help="Лист XLSX (по умолчанию: main_list)"
    )
    parser.add_argument(
        "--notify",
        choices=("all", "delegates", "none"),
        default=None,
        help="Уведомления при назначении прав",
    )
    parser.add_argument("--dry-run", action="store_true", help="Проверить без изменений")
    parser.add_argument(
        "--task-timeout",
        type=float,
        default=60.0,
        help="Ожидание асинхронной задачи в секундах (по умолчанию: 60)",
    )
    parser.add_argument(
        "--report",
        default="shared_mailboxes_result.csv",
        help="Путь к итоговому CSV-отчёту",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        load_dotenv_file()
    except (ValidationError, OSError) as exc:
        die(str(exc))
    token = (os.getenv("OAUTH_TOKEN") or os.getenv("TOKEN") or "").strip()
    org_id = (os.getenv("ORG_ID") or "").strip()
    notify = args.notify or (os.getenv("NOTIFY") or "none").strip().casefold()

    if not token:
        die("задайте OAUTH_TOKEN в .env (также поддерживается прежнее имя TOKEN)")
    if not org_id:
        die("задайте ORG_ID в .env")
    if not org_id.isdigit():
        die("ORG_ID должен состоять из цифр")
    if notify not in {"all", "delegates", "none"}:
        die("NOTIFY должен быть all, delegates или none")

    try:
        plans = load_plan(Path(args.file), args.sheet)
        api = Api360(token, org_id)
        warnings = resolve_actor_ids(plans, api)
    except (ValidationError, OSError, RuntimeError) as exc:
        die(str(exc))

    grant_count = sum(len(plan.grants) for plan in plans)
    print(f"Проверка пройдена: общих ящиков — {len(plans)}, назначений — {grant_count}")
    for warning in warnings:
        print(f"ПРЕДУПРЕЖДЕНИЕ: {warning}")

    for plan in plans:
        print(f"  {plan.email}: {len(plan.grants)} назначений")
        for grant in plan.grants:
            print(
                f"    {grant.identifier} -> UID {grant.actor_id}: {', '.join(grant.roles)}"
            )

    if args.dry_run:
        print("DRY-RUN: запросы на создание и изменение не отправлялись.")
        return 0

    report_rows: list[dict[str, str]] = []
    failures = 0
    for plan in plans:
        try:
            plan.resource_id = api.create_shared_mailbox(plan)
            print(f"СОЗДАН: {plan.email} -> resourceId={plan.resource_id}")
        except Exception as exc:
            failures += len(plan.grants)
            print(f"ОШИБКА СОЗДАНИЯ {plan.email}: {exc}", file=sys.stderr)
            for grant in plan.grants:
                report_rows.append(
                    {
                        "mailbox_email": plan.email,
                        "resource_id": "",
                        "delegate": grant.identifier,
                        "actor_id": grant.actor_id or "",
                        "roles": ",".join(grant.roles),
                        "status": "create_error",
                        "details": str(exc),
                    }
                )
            continue

        for grant in plan.grants:
            try:
                task_id = api.set_access(plan.resource_id, grant, notify)
                status = api.wait_task(task_id, args.task_timeout)
                if status != "complete":
                    failures += 1
                details = f"taskId={task_id}"
                print(
                    f"  {status.upper()}: {grant.identifier} -> UID {grant.actor_id} "
                    f"({details})"
                )
            except Exception as exc:
                failures += 1
                status = "request_error"
                details = str(exc)
                print(f"  ОШИБКА: {grant.identifier}: {exc}", file=sys.stderr)
            report_rows.append(
                {
                    "mailbox_email": plan.email,
                    "resource_id": plan.resource_id,
                    "delegate": grant.identifier,
                    "actor_id": grant.actor_id or "",
                    "roles": ",".join(grant.roles),
                    "status": status,
                    "details": details,
                }
            )

    report_path = Path(args.report)
    write_report(report_path, report_rows)
    print(f"Отчёт: {report_path.resolve()}")
    print(f"Завершено. Ошибок: {failures}")
    return 1 if failures else 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        die("выполнение прервано пользователем", 130)
